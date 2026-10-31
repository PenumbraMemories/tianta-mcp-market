"""PPT Generator MCP：把结构化大纲渲染成可编辑的 PowerPoint。

这个技能只负责排版，不调用任何在线模型或图片服务：
- 输入是主模型生成的一份 deck JSON，输出是本地 .pptx。
- 使用随包附带的 python-pptx / Pillow / lxml，离线运行。
- 版式、配色和模板参考 slide-deck-ai（MIT），但重新实现了更简洁的渲染层。
"""

from __future__ import annotations

import json
import os
import re
import sys
import tempfile
from datetime import date
from pathlib import Path

MODULE_ROOT = Path(__file__).resolve().parent
VENDOR_DIRECTORY = MODULE_ROOT / "vendor"
if VENDOR_DIRECTORY.is_dir():
    sys.path.insert(0, str(VENDOR_DIRECTORY))

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement


PROTOCOL_VERSION = "2024-11-05"
MODULE_ID = "ppt-generator"
MODULE_NAME = "PPT 生成"
MODULE_VERSION = "1.0.0"

SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5
MARGIN = 0.78
CONTENT_WIDTH = 11.77
MAX_SLIDES = 40
MAX_BULLETS = 12

BODY_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"

TEMPLATES = {
    "ion": {
        "file": "Ion_Boardroom.pptx",
        "label": "紫罗兰",
        "title_font": "Microsoft YaHei",
        "bg": "F8F5FB",
        "ink": "24172E",
        "muted": "6E5D7A",
        "accent": "A21C63",
        "accent2": "6D3BB5",
        "card": "FFFFFF",
        "soft": "F0E8F6",
        "line": "E3D9EA",
        "dark": "1E1229",
        "on_dark": "F7F2FA",
    },
    "urban": {
        "file": "Urban_monochrome.pptx",
        "label": "暖灰",
        "title_font": "SimSun",
        "bg": "F4F1E8",
        "ink": "20201D",
        "muted": "6B665A",
        "accent": "B45309",
        "accent2": "3F4A4C",
        "card": "FFFFFF",
        "soft": "EDE7DA",
        "line": "D9D1C0",
        "dark": "1F2424",
        "on_dark": "F5F1E6",
    },
    "minimal": {
        "file": "Minimalist_sales_pitch.pptx",
        "label": "极简",
        "title_font": "Microsoft YaHei",
        "bg": "FFFFFF",
        "ink": "111111",
        "muted": "6B6B6B",
        "accent": "111111",
        "accent2": "C2410C",
        "card": "F6F5F2",
        "soft": "F1EFE9",
        "line": "E3E1DB",
        "dark": "111111",
        "on_dark": "F7F5EF",
    },
    "blank": {
        "file": "Blank.pptx",
        "label": "经典蓝",
        "title_font": "Microsoft YaHei",
        "bg": "F7F9FC",
        "ink": "0F172A",
        "muted": "5B6B82",
        "accent": "2563EB",
        "accent2": "0EA5E9",
        "card": "FFFFFF",
        "soft": "EAF1FB",
        "line": "DCE5F0",
        "dark": "0B1220",
        "on_dark": "F1F5F9",
    },
}

LAYOUTS = {
    "section",
    "bullets",
    "two_column",
    "comparison",
    "image",
    "table",
    "quote",
    "stats",
    "closing",
}

RICH_PATTERN = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\*[^*]+\*)")


def _send(message):
    sys.stdout.write(json.dumps(message, ensure_ascii=False) + "\n")
    sys.stdout.flush()


def _log(message):
    sys.stderr.write(str(message) + "\n")
    sys.stderr.flush()


def _send_result(request_id, result):
    _send({"jsonrpc": "2.0", "id": request_id, "result": result})


def _send_error(request_id, message):
    _send(
        {
            "jsonrpc": "2.0",
            "id": request_id,
            "error": {"code": -32000, "message": str(message)},
        }
    )


def _color(value):
    return RGBColor.from_string(str(value).lstrip("#").upper())


def _mix(first, second, ratio):
    ratio = max(0.0, min(1.0, ratio))
    a = _color(first)
    b = _color(second)
    return RGBColor(
        int(a[0] + (b[0] - a[0]) * ratio),
        int(a[1] + (b[1] - a[1]) * ratio),
        int(a[2] + (b[2] - a[2]) * ratio),
    )


def _clean(value, limit=None):
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if limit and len(text) > limit:
        text = text[: limit - 1].rstrip() + "…"
    return text


def _absolute_path(value, label):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(label + " 必须是绝对路径。")
    path = Path(value.strip()).expanduser()
    if not path.is_absolute():
        raise ValueError(label + " 必须是绝对路径：" + value)
    return path.resolve()


def _style_run(run, size, color, bold=False, italic=False, font=None, spacing=0):
    run.font.name = font or BODY_FONT
    run.font.size = Pt(max(1, float(size)))
    run.font.bold = bool(bold)
    run.font.italic = bool(italic)
    run.font.color.rgb = _color(color)

    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            r_pr.append(node)
        node.set("typeface", font or BODY_FONT)
    if spacing:
        r_pr.set("spc", str(int(spacing * 100)))


def _add_runs(paragraph, text, size, color, bold=False, italic=False, font=None):
    text = str(text or "")
    position = 0
    for match in RICH_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run()
            run.text = text[position : match.start()]
            _style_run(run, size, color, bold, italic, font)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run()
            run.text = token[2:-2]
            _style_run(run, size, color, True, italic, font)
        elif token.startswith("__") and token.endswith("__"):
            run = paragraph.add_run()
            run.text = token[2:-2]
            _style_run(run, size, color, True, italic, font)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run()
            run.text = token[1:-1]
            _style_run(run, size * 0.92, color, False, False, CODE_FONT)
        else:
            run = paragraph.add_run()
            run.text = token[1:-1]
            _style_run(run, size, color, bold, True, font)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run()
        run.text = text[position:]
        _style_run(run, size, color, bold, italic, font)


def _textbox(
    slide,
    text,
    x,
    y,
    width,
    height,
    size,
    color,
    bold=False,
    font=None,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing=1.0,
    margins=(0, 0, 0, 0),
    spacing=0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left = Inches(margins[0])
    frame.margin_top = Inches(margins[1])
    frame.margin_right = Inches(margins[2])
    frame.margin_bottom = Inches(margins[3])

    lines = str(text or "").split("\n")
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    _add_runs(paragraph, lines[0], size, color, bold, False, font)
    if spacing:
        for run in paragraph.runs:
            run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))

    for line in lines[1:]:
        paragraph = frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        _add_runs(paragraph, line, size, color, bold, False, font)
        if spacing:
            for run in paragraph.runs:
                run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
    return shape


def _rect(slide, x, y, width, height, fill, line=None, radius=False, line_width=1):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _color(line)
        shape.line.width = Pt(line_width)
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except (IndexError, ValueError):
            pass
    return shape


def _circle(slide, x, y, diameter, fill=None, line=None, line_width=1.25):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _color(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _color(line)
        shape.line.width = Pt(line_width)
    return shape


def _paint_background(slide, color):
    shape = _rect(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, color)
    # 背景必须位于所有内容之下。
    sp_tree = slide.shapes._spTree
    sp_tree.remove(shape._element)
    sp_tree.insert(2, shape._element)


def _decorate(slide, palette):
    _rect(slide, 0, 0, 0.075, SLIDE_HEIGHT, palette["accent"])
    _rect(slide, 12.12, 0.55, 0.16, 0.16, palette["accent"])


def _remove_bullet(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont", "a:buSzPct"):
        for node in p_pr.findall(qn(tag)):
            p_pr.remove(node)


def _set_bullet(paragraph, char, color, indent=0.30, hanging=0.22):
    p_pr = paragraph._p.get_or_add_pPr()
    _remove_bullet(paragraph)
    p_pr.set("marL", str(int(Inches(indent))))
    p_pr.set("indent", str(-int(Inches(hanging))))

    bullet_color = OxmlElement("a:buClr")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", str(color))
    bullet_color.append(srgb)
    p_pr.append(bullet_color)

    bullet_font = OxmlElement("a:buFont")
    bullet_font.set("typeface", "Arial")
    p_pr.append(bullet_font)

    bullet_char = OxmlElement("a:buChar")
    bullet_char.set("char", char)
    p_pr.append(bullet_char)


def _flatten_bullets(items, level=0):
    result = []
    if isinstance(items, str):
        result.append((items, level))
    elif isinstance(items, dict):
        text = items.get("text") or items.get("title") or items.get("label")
        if text:
            result.append((str(text), int(items.get("level", level))))
    elif isinstance(items, (list, tuple)):
        for item in items:
            result.extend(_flatten_bullets(item, level + 1 if isinstance(item, list) else level))
    return result


def _add_bullets(
    slide,
    items,
    x,
    y,
    width,
    height,
    size,
    palette,
    numbered=False,
    color=None,
    anchor=MSO_ANCHOR.TOP,
):
    flat = _flatten_bullets(items)
    flat = [(text, level) for text, level in flat if _clean(text)]
    flat = flat[:MAX_BULLETS]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0

    text_color = color or palette["ink"]
    for index, (text, level) in enumerate(flat):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = 1.10
        paragraph.space_after = Pt(8 if level == 0 else 3)
        current_size = size if level == 0 else max(10, size * 0.84)
        current_color = text_color if level == 0 else palette["muted"]

        if numbered and level == 0:
            prefix = paragraph.add_run()
            prefix.text = f"{index + 1:02d}  "
            _style_run(prefix, current_size, palette["accent"], True)
            _add_runs(paragraph, text, current_size, current_color)
            p_pr = paragraph._p.get_or_add_pPr()
            p_pr.set("marL", str(int(Inches(0.52))))
            p_pr.set("indent", str(-int(Inches(0.52))))
        else:
            char = "•" if level == 0 else ("–" if level == 1 else "·")
            _set_bullet(
                paragraph,
                char,
                palette["accent"] if level == 0 else palette["muted"],
                0.30 if level == 0 else 0.62,
                0.22,
            )
            _add_runs(paragraph, text, current_size, current_color)
    return box


def _header(slide, title, eyebrow, palette):
    title = _clean(title, 120) or "未命名页面"
    if eyebrow:
        _textbox(
            slide,
            _clean(eyebrow, 50),
            MARGIN,
            0.46,
            9.4,
            0.28,
            9,
            palette["accent"],
            True,
            palette["title_font"],
            spacing=1.4,
        )
    title_size = 30
    if len(title) > 46:
        title_size = 21
    elif len(title) > 34:
        title_size = 24
    elif len(title) > 27:
        title_size = 27
    _textbox(
        slide,
        title,
        MARGIN,
        0.76,
        CONTENT_WIDTH,
        0.96,
        title_size,
        palette["ink"],
        True,
        palette["title_font"],
        anchor=MSO_ANCHOR.TOP,
        line_spacing=1.0,
    )
    _rect(slide, MARGIN, 1.74, 0.72, 0.055, palette["accent"])
    _rect(slide, MARGIN, 1.785, CONTENT_WIDTH, 0.012, palette["line"])


def _footer(slide, deck, page, total, palette, dark=False):
    color = _mix(palette["on_dark"], palette["dark"], 0.55) if dark else palette["muted"]
    line = _mix(palette["on_dark"], palette["dark"], 0.28) if dark else palette["line"]
    _rect(slide, MARGIN, 6.98, CONTENT_WIDTH, 0.01, line)
    footer = _clean(deck.get("footer") or deck.get("title"), 60)
    _textbox(slide, footer, MARGIN, 7.07, 8.4, 0.22, 8, color, font=BODY_FONT)
    _textbox(
        slide,
        f"{page:02d} / {total:02d}",
        10.55,
        7.07,
        2.0,
        0.22,
        8,
        color,
        font=BODY_FONT,
        align=PP_ALIGN.RIGHT,
    )


def _key_card(slide, text, palette, y=5.92, height=0.78):
    text = _clean(text, 180)
    if not text:
        return
    _rect(slide, MARGIN, y, CONTENT_WIDTH, height, _mix(palette["bg"], palette["accent"], 0.08))
    _rect(slide, MARGIN, y, 0.06, height, palette["accent"])
    _textbox(
        slide,
        "KEY TAKEAWAY",
        MARGIN + 0.22,
        y + 0.12,
        2.2,
        0.2,
        7.8,
        palette["accent"],
        True,
        BODY_FONT,
        spacing=1.5,
    )
    size = 13.5 if len(text) <= 62 else 11.5
    _textbox(
        slide,
        text,
        MARGIN + 2.05,
        y + 0.10,
        CONTENT_WIDTH - 2.3,
        height - 0.18,
        size,
        palette["ink"],
        font=BODY_FONT,
        anchor=MSO_ANCHOR.MIDDLE,
        line_spacing=1.05,
    )


def _cover(prs, deck, palette):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _paint_background(slide, palette["bg"])
    _decorate(slide, palette)

    cover_image = deck.get("cover_image")
    if cover_image:
        path = Path(str(cover_image)).expanduser()
        if path.is_file():
            _add_picture_cover(slide, path, 7.5, 0, 5.83, SLIDE_HEIGHT)
            _rect(slide, 7.5, 0, 0.06, SLIDE_HEIGHT, palette["accent"])
        else:
            _circle(slide, 9.3, 1.1, 4.6, _mix(palette["bg"], palette["accent"], 0.12))
    else:
        _circle(slide, 9.25, 0.9, 4.9, _mix(palette["bg"], palette["accent"], 0.11))
        _circle(slide, 10.9, 2.65, 2.1, None, _mix(palette["bg"], palette["accent"], 0.55), 1.2)
        _circle(slide, 11.65, 1.25, 0.22, palette["accent"])

    _textbox(
        slide,
        _clean(deck.get("eyebrow") or "PRESENTATION", 40),
        0.95,
        1.05,
        7.0,
        0.3,
        9.5,
        palette["accent"],
        True,
        BODY_FONT,
        spacing=1.8,
    )
    title = _clean(deck.get("title"), 90) or "未命名演示"
    text_width = 6.15 if cover_image else 7.0
    title_size = 46
    if len(title) > 34:
        title_size = 34
    elif len(title) > 24:
        title_size = 39
    _textbox(
        slide,
        title,
        0.95,
        1.48,
        text_width if cover_image else 8.0,
        2.2,
        title_size,
        palette["ink"],
        True,
        palette["title_font"],
        line_spacing=1.04,
    )
    _rect(slide, 0.95, 3.92, 1.15, 0.065, palette["accent"])
    subtitle = _clean(deck.get("subtitle"), 120)
    if subtitle:
        _textbox(
            slide,
            subtitle,
            0.95,
            4.18,
            text_width if cover_image else 8.2,
            1.1,
            17,
            palette["muted"],
            font=BODY_FONT,
            line_spacing=1.25,
        )

    meta = []
    if _clean(deck.get("author")):
        meta.append(_clean(deck.get("author"), 50))
    meta.append(_clean(deck.get("date") or date.today().strftime("%Y.%m"), 30))
    _textbox(
        slide,
        "    ·    ".join(meta),
        0.95,
        6.28,
        text_width if cover_image else 7.6,
        0.3,
        9.5,
        palette["muted"],
        font=BODY_FONT,
        spacing=0.8,
    )
    return slide


def _section(prs, slide_data, index, palette):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _paint_background(slide, palette["bg"])
    _decorate(slide, palette)

    number = f"{index:02d}"
    _textbox(
        slide,
        number,
        9.5,
        1.15,
        3.2,
        2.8,
        150,
        _mix(palette["bg"], palette["accent"], 0.16),
        True,
        palette["title_font"],
        align=PP_ALIGN.RIGHT,
    )
    _textbox(
        slide,
        _clean(slide_data.get("eyebrow") or "SECTION", 40),
        0.95,
        2.0,
        7.2,
        0.3,
        9,
        palette["accent"],
        True,
        BODY_FONT,
        spacing=1.8,
    )
    title = _clean(slide_data.get("title"), 80) or "章节"
    _textbox(
        slide,
        title,
        0.95,
        2.35,
        8.6,
        1.8,
        40 if len(title) <= 20 else 32,
        palette["ink"],
        True,
        palette["title_font"],
        line_spacing=1.05,
    )
    _rect(slide, 0.95, 4.42, 1.1, 0.06, palette["accent"])
    subtitle = _clean(slide_data.get("subtitle"), 100)
    if subtitle:
        _textbox(
            slide,
            subtitle,
            0.95,
            4.68,
            8.0,
            1.0,
            15,
            palette["muted"],
            font=BODY_FONT,
            line_spacing=1.2,
        )
    return slide


def _bullets(prs, slide_data, index, palette):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _paint_background(slide, palette["bg"])
    _decorate(slide, palette)
    _header(
        slide,
        slide_data.get("title"),
        slide_data.get("eyebrow") or f"{index:02d}",
        palette,
    )

    items = slide_data.get("bullets") or slide_data.get("points") or []
    flat = _flatten_bullets(items)
    key_message = _clean(slide_data.get("key_message"), 180)
    body_height = 3.55 if key_message else 4.24
    body_size = 18
    if len(flat) > 6:
        body_size = 14
    elif len(flat) > 4:
        body_size = 15.5

    if len(flat) > 6:
        top_level = [position for position, (_, level) in enumerate(flat) if level == 0]
        target = len(flat) / 2
        middle = min(top_level, key=lambda position: abs(position - target)) if top_level else (len(flat) + 1) // 2
        if middle == 0 and len(top_level) > 1:
            middle = top_level[1]
        if middle >= len(flat) and len(top_level) > 1:
            middle = top_level[-2]
        left_items = [{"text": text, "level": level} for text, level in flat[:middle]]
        right_items = [{"text": text, "level": level} for text, level in flat[middle:]]
        _add_bullets(
            slide, left_items, MARGIN, 2.0, 5.5, body_height, 14, palette
        )
        _add_bullets(
            slide,
            right_items,
            MARGIN + 6.1,
            2.0,
            5.5,
            body_height,
            14,
            palette,
        )
    else:
        _add_bullets(
            slide,
            items,
            MARGIN,
            2.02,
            10.9,
            body_height,
            body_size,
            palette,
            numbered=bool(slide_data.get("numbered")),
            anchor=MSO_ANCHOR.MIDDLE,
        )
    if key_message:
        _key_card(slide, key_message, palette)
    return slide


def _two_column(prs, slide_data, index, palette, comparison=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _paint_background(slide, palette["bg"])
    _decorate(slide, palette)
    _header(
        slide,
        slide_data.get("title"),
        slide_data.get("eyebrow") or f"{index:02d}",
        palette,
    )

    columns = slide_data.get("columns") or []
    if not isinstance(columns, list) or len(columns) < 2:
        columns = [
            {"heading": "现状", "bullets": slide_data.get("bullets") or []},
            {"heading": "目标", "bullets": slide_data.get("bullets_2") or []},
        ]
    columns = columns[:3]
    gap = 0.48
    width = (CONTENT_WIDTH - gap * (len(columns) - 1)) / len(columns)
    key_message = _clean(slide_data.get("key_message"), 180)
    card_height = 3.55 if key_message else 4.3

    for column_index, column in enumerate(columns):
        x = MARGIN + column_index * (width + gap)
        _rect(slide, x, 2.02, width, card_height, palette["card"])
        _rect(slide, x, 2.02, width, 0.055, palette["accent"])
        heading = _clean(column.get("heading") or column.get("title"), 40)
        if heading:
            _textbox(
                slide,
                heading,
                x + 0.24,
                2.22,
                width - 0.48,
                0.38,
                16 if len(columns) == 2 else 14,
                palette["accent"],
                True,
                palette["title_font"],
            )
        _add_bullets(
            slide,
            column.get("bullets") or column.get("points") or [],
            x + 0.24,
            2.72,
            width - 0.48,
            card_height - 0.85,
            13.5 if len(columns) == 2 else 12,
            palette,
        )
        if comparison and len(columns) == 2 and column_index == 0:
            _circle(
                slide,
                MARGIN + width + gap / 2 - 0.28,
                2.02 + card_height / 2 - 0.28,
                0.56,
                palette["bg"],
                palette["accent"],
                1.2,
            )
            _textbox(
                slide,
                "VS",
                MARGIN + width + gap / 2 - 0.28,
                2.02 + card_height / 2 - 0.10,
                0.56,
                0.2,
                8,
                palette["accent"],
                True,
                BODY_FONT,
                align=PP_ALIGN.CENTER,
            )
    if key_message:
        _key_card(slide, key_message, palette)
    return slide


def _add_picture_cover(slide, path, x, y, width, height):
    try:
        with Image.open(path) as image:
            image_width, image_height = image.size
    except Exception as error:
        raise ValueError("无法读取图片：" + str(error)) from error
    if image_width <= 0 or image_height <= 0:
        raise ValueError("图片尺寸无效：" + str(path))
    picture = slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame_ratio = width / height
    image_ratio = image_width / image_height
    if image_ratio > frame_ratio:
        crop = (1 - frame_ratio / image_ratio) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    else:
        crop = (1 - image_ratio / frame_ratio) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def _image(prs, slide_data, index, palette, warnings):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _paint_background(slide, palette["bg"])
    _decorate(slide, palette)
    _header(
        slide,
        slide_data.get("title"),
        slide_data.get("eyebrow") or f"{index:02d}",
        palette,
    )

    image_value = slide_data.get("image") or slide_data.get("image_path") or ""
    image_path = Path(str(image_value)).expanduser() if image_value else None
    image_on_left = str(slide_data.get("image_position", "right")).lower() == "left"
    print_width = 5.35
    print_left = MARGIN if image_on_left else 7.25
    text_left = 7.25 if image_on_left else MARGIN
    text_width = 5.15

    _add_bullets(
        slide,
        slide_data.get("bullets") or slide_data.get("points") or [],
        text_left,
        2.02,
        text_width,
        3.55 if slide_data.get("key_message") else 4.2,
        15,
        palette,
    )
    if slide_data.get("key_message"):
        _key_card(slide, slide_data.get("key_message"), palette)

    box_top = 2.02
    box_height = 3.78
    if image_path and image_path.is_file():
        try:
            _add_picture_cover(
                slide, image_path, print_left, box_top, print_width, box_height
            )
        except Exception as error:
            warnings.append(str(error))
            _image_placeholder(slide, print_left, box_top, print_width, box_height, palette)
    else:
        if image_value:
            warnings.append("图片不存在，已使用占位框：" + str(image_value))
        _image_placeholder(slide, print_left, box_top, print_width, box_height, palette)

    caption = _clean(slide_data.get("image_caption"), 90)
    if caption:
        _textbox(
            slide,
            caption,
            print_left,
            box_top + box_height + 0.12,
            print_width,
            0.3,
            9,
            palette["muted"],
            font=BODY_FONT,
            align=PP_ALIGN.CENTER,
        )
    return slide


def _image_placeholder(slide, x, y, width, height, palette):
    _rect(slide, x, y, width, height, _mix(palette["bg"], palette["accent"], 0.06))
    _circle(slide, x + width / 2 - 0.38, y + height / 2 - 0.62, 0.76, None, palette["line"], 1.2)
    _textbox(
        slide,
        "图片",
        x,
        y + height / 2 + 0.15,
        width,
        0.3,
        10,
        palette["muted"],
        font=BODY_FONT,
        align=PP_ALIGN.CENTER,
    )


def _table(prs, slide_data, index, palette, warnings):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _paint_background(slide, palette["bg"])
    _decorate(slide, palette)
    _header(
        slide,
        slide_data.get("title"),
        slide_data.get("eyebrow") or f"{index:02d}",
        palette,
    )

    table_data = slide_data.get("table") or {}
    headers = [str(value) for value in (table_data.get("headers") or [])]
    rows = [[str(value) for value in row] for row in (table_data.get("rows") or [])]
    if not headers or not rows:
        warnings.append("表格页缺少 headers 或 rows，已渲染为空白提示。")
        _textbox(
            slide,
            "未提供表格数据",
            MARGIN,
            3.0,
            CONTENT_WIDTH,
            0.6,
            18,
            palette["muted"],
            align=PP_ALIGN.CENTER,
        )
        return slide

    column_count = min(len(headers), max(len(row) for row in rows))
    headers = headers[:column_count]
    rows = [row[:column_count] for row in rows]
    if len(rows) > 8:
        warnings.append("表格超过 8 行，已压缩字号，建议拆分页面。")
    row_count = len(rows) + 1
    font_size = 11.5 if row_count <= 6 else (10.5 if row_count <= 8 else 9.5)
    row_height = 0.44 if row_count <= 6 else (0.38 if row_count <= 8 else 0.32)
    table_height = min(4.45, row_height * row_count)

    weights = []
    for column_index in range(column_count):
        values = [headers[column_index]] + [row[column_index] for row in rows]
        weight = max(
            6,
            max(
                sum(1.85 if ord(char) > 127 else 1.0 for char in value)
                for value in values
            ),
        )
        weights.append(min(weight, 34))
    total_weight = sum(weights)
    widths = [max(1.15, CONTENT_WIDTH * weight / total_weight) for weight in weights]
    scale = CONTENT_WIDTH / sum(widths)
    widths = [width * scale for width in widths]

    table_top = 2.0
    table_shape = slide.shapes.add_table(
        row_count,
        column_count,
        Inches(MARGIN),
        Inches(table_top),
        Inches(CONTENT_WIDTH),
        Inches(table_height),
    )
    table = table_shape.table
    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Inches(width)
    for row_index in range(row_count):
        table.rows[row_index].height = Inches(table_height / row_count)

    for row_index in range(row_count):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.margin_left = Inches(0.10)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            value = headers[column_index] if row_index == 0 else rows[row_index - 1][column_index]
            cell.text = value
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.name = BODY_FONT
                run.font.size = Pt(font_size)
                run.font.bold = row_index == 0
                run.font.color.rgb = _color(
                    palette["on_dark"] if row_index == 0 else palette["ink"]
                )
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = _color(palette["accent"])
            else:
                cell.fill.fore_color.rgb = _color(
                    palette["card"] if row_index % 2 == 1 else palette["soft"]
                )

    if slide_data.get("key_message"):
        _key_card(slide, slide_data.get("key_message"), palette, y=6.02, height=0.72)
    return slide


def _quote(prs, slide_data, page, total, palette, deck):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _paint_background(slide, palette["dark"])
    _rect(slide, 0, 0, 0.075, SLIDE_HEIGHT, palette["accent"])
    _textbox(
        slide,
        "“",
        0.72,
        0.55,
        1.5,
        1.4,
        110,
        _mix(palette["dark"], palette["accent"], 0.55),
        True,
        "Georgia",
    )
    quote = _clean(slide_data.get("quote") or slide_data.get("title"), 260)
    _textbox(
        slide,
        quote,
        1.35,
        2.05,
        10.6,
        2.8,
        30 if len(quote) <= 70 else 24,
        palette["on_dark"],
        False,
        palette["title_font"],
        line_spacing=1.22,
    )
    _rect(slide, 1.35, 5.12, 0.9, 0.045, palette["accent"])
    attribution = _clean(slide_data.get("attribution") or slide_data.get("subtitle"), 90)
    if attribution:
        _textbox(
            slide,
            attribution,
            1.35,
            5.3,
            8.5,
            0.4,
            12,
            _mix(palette["on_dark"], palette["dark"], 0.35),
            True,
            BODY_FONT,
            spacing=1.0,
        )
    _footer(slide, deck, page, total, palette, dark=True)
    return slide


def _stats(prs, slide_data, index, palette):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _paint_background(slide, palette["bg"])
    _decorate(slide, palette)
    _header(
        slide,
        slide_data.get("title"),
        slide_data.get("eyebrow") or f"{index:02d}",
        palette,
    )

    metrics = slide_data.get("metrics") or slide_data.get("stats") or []
    metrics = metrics[:4]
    if not metrics:
        metrics = [{"value": "—", "label": "暂无数据", "note": ""}]
    gap = 0.28
    width = (CONTENT_WIDTH - gap * (len(metrics) - 1)) / len(metrics)
    key_message = _clean(slide_data.get("key_message"), 180)
    height = 3.7 if key_message else 4.4
    for metric_index, metric in enumerate(metrics):
        x = MARGIN + metric_index * (width + gap)
        _rect(slide, x, 2.02, width, height, palette["card"])
        _rect(slide, x, 2.02, width, 0.055, palette["accent"])
        value = _clean(metric.get("value"), 20) or "—"
        _textbox(
            slide,
            value,
            x + 0.22,
            2.32,
            width - 0.44,
            0.8,
            34 if len(value) <= 8 else 25,
            palette["accent"],
            True,
            palette["title_font"],
        )
        _textbox(
            slide,
            _clean(metric.get("label"), 40),
            x + 0.22,
            3.22,
            width - 0.44,
            0.5,
            13,
            palette["ink"],
            True,
            BODY_FONT,
        )
        note = _clean(metric.get("note"), 90)
        if note:
            _textbox(
                slide,
                note,
                x + 0.22,
                3.78,
                width - 0.44,
                height - 1.98,
                10.5,
                palette["muted"],
                font=BODY_FONT,
                line_spacing=1.15,
            )
    if key_message:
        _key_card(slide, key_message, palette)
    return slide


def _closing(prs, deck, page, total, palette):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _paint_background(slide, palette["dark"])
    _rect(slide, 0, 0, 0.075, SLIDE_HEIGHT, palette["accent"])
    _circle(slide, 9.7, -1.0, 5.0, _mix(palette["dark"], palette["accent"], 0.14))
    _circle(slide, 10.9, 4.5, 3.1, None, _mix(palette["dark"], palette["accent"], 0.35), 1.0)

    closing = deck.get("closing") if isinstance(deck.get("closing"), dict) else {}
    title = _clean(closing.get("title") or "谢谢观看", 40)
    subtitle = _clean(
        closing.get("subtitle") or "欢迎交流，也欢迎把下一步交给我。",
        100,
    )
    _textbox(
        slide,
        "THANK YOU",
        1.05,
        1.45,
        5.2,
        0.3,
        9,
        palette["accent"],
        True,
        BODY_FONT,
        spacing=1.8,
    )
    _textbox(
        slide,
        title,
        1.05,
        1.9,
        8.4,
        1.4,
        48 if len(title) <= 12 else 36,
        palette["on_dark"],
        True,
        palette["title_font"],
    )
    _rect(slide, 1.05, 3.55, 1.1, 0.06, palette["accent"])
    _textbox(
        slide,
        subtitle,
        1.05,
        3.82,
        8.0,
        1.1,
        16,
        _mix(palette["on_dark"], palette["dark"], 0.3),
        font=BODY_FONT,
        line_spacing=1.25,
    )
    footer = _clean(deck.get("footer") or deck.get("title"), 60)
    _textbox(
        slide,
        footer,
        1.05,
        6.28,
        8.0,
        0.3,
        9,
        _mix(palette["on_dark"], palette["dark"], 0.5),
        font=BODY_FONT,
    )
    _textbox(
        slide,
        f"{page:02d} / {total:02d}",
        10.55,
        6.28,
        2.0,
        0.3,
        8,
        _mix(palette["on_dark"], palette["dark"], 0.5),
        font=BODY_FONT,
        align=PP_ALIGN.RIGHT,
    )
    return slide


def _layout_for(slide_data):
    layout = str(slide_data.get("layout") or "").lower().strip()
    if layout:
        if layout not in LAYOUTS:
            raise ValueError(
                "不支持的 layout："
                + layout
                + "；可用："
                + ", ".join(sorted(LAYOUTS))
            )
        return layout
    if slide_data.get("table"):
        return "table"
    if slide_data.get("quote"):
        return "quote"
    if slide_data.get("columns"):
        return "two_column"
    if slide_data.get("image") or slide_data.get("image_path"):
        return "image"
    if slide_data.get("metrics") or slide_data.get("stats"):
        return "stats"
    return "bullets"


def _validate_deck(deck):
    if not isinstance(deck, dict):
        raise ValueError("deck 必须是 JSON 对象。")
    title = _clean(deck.get("title"), 90)
    if not title:
        raise ValueError("deck.title 不能为空。")
    slides = deck.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("deck.slides 必须是非空数组。")
    if len(slides) > MAX_SLIDES:
        raise ValueError(f"单次最多生成 {MAX_SLIDES} 页，当前 {len(slides)} 页。")
    for index, slide_data in enumerate(slides, start=1):
        if not isinstance(slide_data, dict):
            raise ValueError(f"第 {index} 页必须是对象。")
        _layout_for(slide_data)
    return deck


def _template_path(template):
    key = str(template or "urban").lower().strip()
    if key not in TEMPLATES:
        raise ValueError(
            "未知模板：" + str(template) + "；可用：" + ", ".join(TEMPLATES.keys())
        )
    path = MODULE_ROOT / "templates" / TEMPLATES[key]["file"]
    if not path.is_file():
        raise RuntimeError("模板文件缺失：" + str(path))
    return key, path


def create_presentation(arguments):
    deck = arguments.get("deck")
    if isinstance(deck, str):
        try:
            deck = json.loads(deck)
        except ValueError as error:
            raise ValueError("deck 不是有效 JSON：" + str(error)) from error
    _validate_deck(deck)

    template, template_path = _template_path(arguments.get("template"))
    palette = TEMPLATES[template]
    output_path = _absolute_path(arguments.get("output_path"), "output_path")
    if output_path.suffix.lower() != ".pptx":
        raise ValueError("output_path 必须以 .pptx 结尾。")
    overwrite = bool(arguments.get("overwrite", False))
    if output_path.exists() and not overwrite:
        raise ValueError("输出文件已存在；如需覆盖请设置 overwrite=true：" + str(output_path))
    output_path.parent.mkdir(parents=True, exist_ok=True)

    warnings = []
    presentation = Presentation(str(template_path))
    slides = deck["slides"]
    has_closing = _layout_for(slides[-1]) == "closing"
    closing = bool(deck.get("closing")) and not has_closing
    total = 1 + len(slides) + (1 if closing else 0)

    _cover(presentation, deck, palette)

    for index, slide_data in enumerate(slides, start=1):
        layout = _layout_for(slide_data)
        page = index + 1
        if layout == "section":
            _section(presentation, slide_data, index, palette)
        elif layout == "two_column":
            _two_column(presentation, slide_data, index, palette)
        elif layout == "comparison":
            _two_column(presentation, slide_data, index, palette, comparison=True)
        elif layout == "image":
            _image(presentation, slide_data, index, palette, warnings)
        elif layout == "table":
            _table(presentation, slide_data, index, palette, warnings)
        elif layout == "quote":
            _quote(presentation, slide_data, page, total, palette, deck)
        elif layout == "stats":
            _stats(presentation, slide_data, index, palette)
        elif layout == "closing":
            _closing(presentation, deck, page, total, palette)
        else:
            _bullets(presentation, slide_data, index, palette)

        notes = _clean(
            slide_data.get("notes") or slide_data.get("speaker_notes"), 1200
        )
        if notes:
            presentation.slides[-1].notes_slide.notes_text_frame.text = notes

        if layout not in {"quote", "closing"}:
            _footer(presentation.slides[-1], deck, page, total, palette)

    if closing:
        page = total
        _closing(presentation, deck, page, total, palette)

    if len(presentation.slides) != total:
        raise RuntimeError(
            f"页数校验失败：预期 {total}，实际 {len(presentation.slides)}。"
        )

    file_descriptor, temporary_name = tempfile.mkstemp(
        suffix=".pptx", prefix="tianta-ppt-", dir=str(output_path.parent)
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        presentation.save(str(temporary_path))
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink(missing_ok=True)

    return {
        "ok": True,
        "output_path": str(output_path),
        "template": template,
        "template_label": palette["label"],
        "slide_count": total,
        "file_size": output_path.stat().st_size,
        "warnings": warnings,
    }


TOOL_SCHEMA = {
    "type": "object",
    "properties": {
        "deck": {
            "type": "object",
            "description": (
                "演示内容。title/subtitle/author/date 生成封面；slides 是正文页。"
                "文本支持 **粗体**、*斜体* 和 `等宽`。"
            ),
            "properties": {
                "title": {"type": "string", "description": "封面主标题，也是默认页脚。"},
                "subtitle": {"type": "string", "description": "封面副标题。"},
                "author": {"type": "string", "description": "作者或团队，可省略。"},
                "date": {"type": "string", "description": "日期文本，可省略，默认当前年月。"},
                "eyebrow": {"type": "string", "description": "封面左上小标签，默认 PRESENTATION。"},
                "footer": {"type": "string", "description": "页脚文字，默认使用 title。"},
                "cover_image": {"type": "string", "description": "封面右侧图片的绝对路径，可省略。"},
                "closing": {
                    "type": "object",
                    "description": "自动追加结尾页；若最后一页 layout=closing 则忽略。",
                    "properties": {
                        "title": {"type": "string"},
                        "subtitle": {"type": "string"},
                    },
                    "additionalProperties": False,
                },
                "slides": {
                    "type": "array",
                    "minItems": 1,
                    "maxItems": MAX_SLIDES,
                    "items": {
                        "type": "object",
                        "properties": {
                            "layout": {
                                "type": "string",
                                "enum": [
                                    "section",
                                    "bullets",
                                    "two_column",
                                    "comparison",
                                    "image",
                                    "table",
                                    "quote",
                                    "stats",
                                    "closing",
                                ],
                                "description": "版式；省略时根据字段自动推断。",
                            },
                            "title": {"type": "string", "description": "页面标题。"},
                            "subtitle": {"type": "string", "description": "章节/结尾页副标题。"},
                            "eyebrow": {"type": "string", "description": "标题上方小标签，默认页码。"},
                            "bullets": {
                                "type": "array",
                                "description": "要点数组；元素可以是字符串、{text,level} 或嵌套数组。",
                                "items": {},
                            },
                            "numbered": {
                                "type": "boolean",
                                "description": "bullets 是否显示 01/02 编号，适合目录页。",
                            },
                            "key_message": {"type": "string", "description": "底部一句话结论。"},
                            "columns": {
                                "type": "array",
                                "description": "两栏/三栏内容，每项 {heading, bullets}。",
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "heading": {"type": "string"},
                                        "bullets": {"type": "array", "items": {}},
                                    },
                                    "required": ["heading"],
                                },
                            },
                            "image": {"type": "string", "description": "图片绝对路径。"},
                            "image_position": {
                                "type": "string",
                                "enum": ["left", "right"],
                                "description": "图片在左还是右，默认 right。",
                            },
                            "image_caption": {"type": "string", "description": "图片说明。"},
                            "table": {
                                "type": "object",
                                "properties": {
                                    "headers": {
                                        "type": "array",
                                        "items": {"type": "string"},
                                    },
                                    "rows": {
                                        "type": "array",
                                        "items": {
                                            "type": "array",
                                            "items": {"type": "string"},
                                        },
                                    },
                                },
                                "required": ["headers", "rows"],
                            },
                            "quote": {"type": "string", "description": "quote 页的引文。"},
                            "attribution": {"type": "string", "description": "引文出处。"},
                            "metrics": {
                                "type": "array",
                                "description": "stats 页的数据卡，每项 {value,label,note}。",
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "value": {"type": "string"},
                                        "label": {"type": "string"},
                                        "note": {"type": "string"},
                                    },
                                    "required": ["value", "label"],
                                },
                            },
                            "notes": {"type": "string", "description": "演讲者备注。"},
                        },
                        "additionalProperties": True,
                    },
                },
            },
            "required": ["title", "slides"],
            "additionalProperties": True,
        },
        "output_path": {
            "type": "string",
            "description": "输出 .pptx 的绝对路径；父目录会自动创建。",
        },
        "template": {
            "type": "string",
            "enum": list(TEMPLATES.keys()),
            "default": "urban",
            "description": "视觉模板：urban 暖灰、ion 紫罗兰、minimal 极简、blank 经典蓝。",
        },
        "overwrite": {
            "type": "boolean",
            "default": False,
            "description": "输出文件存在时是否覆盖。",
        },
    },
    "required": ["deck", "output_path"],
    "additionalProperties": False,
}

TOOLS = [
    {
        "name": "create_presentation",
        "description": (
            "把结构化 deck JSON 渲染成可编辑的 .pptx。"
            "deck.title 自动生成封面；slides 支持 section、bullets、two_column、comparison、"
            "image、table、quote、stats、closing 版式。文本支持 **粗体** 和 *斜体*；"
            "图片使用本地绝对路径；表格使用 headers/rows。生成后返回文件路径、页数和警告。"
        ),
        "inputSchema": TOOL_SCHEMA,
    }
]


def _call_tool(name, arguments):
    if name != "create_presentation":
        raise ValueError("未知工具：" + str(name))
    return create_presentation(arguments or {})


def _handle(message):
    method = message.get("method")
    request_id = message.get("id")

    if method == "initialize":
        _send_result(
            request_id,
            {
                "protocolVersion": PROTOCOL_VERSION,
                "capabilities": {"tools": {}},
                "serverInfo": {"name": MODULE_ID, "version": MODULE_VERSION},
            },
        )
    elif method == "notifications/initialized":
        return
    elif method == "ping":
        _send_result(request_id, {})
    elif method == "tools/list":
        _send_result(request_id, {"tools": TOOLS})
    elif method == "tools/call":
        params = message.get("params") or {}
        try:
            result = _call_tool(params.get("name"), params.get("arguments") or {})
            _send_result(
                request_id,
                {
                    "content": [
                        {"type": "text", "text": json.dumps(result, ensure_ascii=False)}
                    ],
                    "isError": False,
                },
            )
        except Exception as error:
            _send_error(request_id, str(error))
    elif request_id is not None:
        _send_error(request_id, "不支持的方法：" + str(method))


def main():
    try:
        sys.stdin.reconfigure(encoding="utf-8")
        sys.stdout.reconfigure(encoding="utf-8", newline="\n")
        sys.stderr.reconfigure(encoding="utf-8")
    except (AttributeError, OSError):
        pass

    for line in sys.stdin:
        text = line.strip()
        if not text:
            continue
        try:
            message = json.loads(text)
        except ValueError:
            continue
        _handle(message)


if __name__ == "__main__":
    main()
